from django.shortcuts import render
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser,FormParser
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework import status
from django.core.files.uploadedfile import InMemoryUploadedFile
from io import BytesIO
from docx import Document
from .serializers import GenerateExamSerializer,CorrectExamSerializer,StudentMarkSerializer,CustomUserSerializer
from rest_framework import generics, status
from .models import QuestionDetail, StudentReview
from django.http import HttpResponse
import csv
from django.contrib.auth.models import User
from rest_framework.permissions import IsAdminUser
from openai import OpenAI
import time
from .utils import extract_details_from_string,save_review_to_db


# Initialize OpenAI client
client = OpenAI(api_key="sk-proj-GmtcP1oqpnaWh0lEuuccT3BlbkFJDAqSxmdbrItZcFJGQwiP")
assistant1 = "asst_knueiNbtyqAK8ZmJUuVbHvPs"
assistant2 = "asst_MZo5dFRj2mO4SaOJTfhexy1v"

# Create your views here.

class GenerateExamAPIView(APIView):
    parser_classes = (MultiPartParser, FormParser)
    serializer_class = GenerateExamSerializer
    def post(self, request, *args, **kwargs):
        serializer = self.serializer_class(data=request.data)
        
        if serializer.is_valid():
            uploaded_file = serializer.validated_data['file']
            exam_chapters = serializer.validated_data['exam_chapters']
            exam_parts = serializer.validated_data['exam_parts']
            each_part_question_numbers = serializer.validated_data['each_part_question_numbers']
            exam_difficulty = serializer.validated_data['exam_difficulty']
            exam_level = serializer.validated_data['exam_level']

            # Read the uploaded file as bytes
            uploaded_file.seek(0)
            file_bytes = uploaded_file.read()
            # Upload the user-provided file to OpenAI
            message_file = client.files.create(
                file=(uploaded_file.name, file_bytes), purpose="assistants"
            )


            # Create a thread
            thread = client.beta.threads.create(
                messages=[
                    {
                        "role": "user",
                        "content": f"Please generate an exam based on the book. Exam should contain the following chapters : {exam_chapters}, The exam should contain the following parts  which are {', '.join(exam_parts)},  each part will have {each_part_question_numbers}. The exam level is {exam_level} and the difficulty is {exam_difficulty}. Your output response should only be the exam text, without any extra instructions or answers.",
                        "attachments": [
                            {"file_id": message_file.id, "tools": [{"type": "file_search"}]}
                        ],
                    }
                ]
            )

            # Start run
            run = client.beta.threads.runs.create_and_poll(
                thread_id=thread.id, assistant_id=assistant1
            )

            # Outputting message
            messages = list(client.beta.threads.messages.list(thread_id=thread.id, run_id=run.id))
            message_content = messages[0].content[0].text

            # return Response(
            #     {'message': message_content},
            #     status=status.HTTP_200_OK
            # )

            # Generate DOCX file
            doc = Document()
            doc.add_heading('Generated Exam', level=1)
            doc.add_paragraph(message_content.value)
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            

            # Create an in-memory file
            in_memory_file = InMemoryUploadedFile(
                buffer, None, 'generated_exam.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', buffer.tell(), None
            )

            response = Response(
                {'message': message_content},
                status=status.HTTP_200_OK
            )
            response['Content-Disposition'] = f'attachment; filename=generated_exam.docx'
            response['Content-Type'] = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            response.content = in_memory_file.read()

            return response

        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
    
class CorrectExamAPIView(APIView):
    def post(self, request, *args, **kwargs):
        serializer = CorrectExamSerializer(data=request.data)
        if serializer.is_valid():
            original_file = serializer.validated_data['original_file']
            student_file = serializer.validated_data['student_file']

            # Ensure both files have supported extensions
            for file in [original_file, student_file]:
                if not file.name.endswith(('.pdf', '.txt')):
                    return Response({'error': f'Unsupported file extension for {file.name}'}, status=status.HTTP_400_BAD_REQUEST)

            vector_store = client.beta.vector_stores.create(name="Book and student paper vector store")

            # Read the file content correctly and prepare for upload
            original_file_stream = original_file.read()
            student_file_stream = student_file.read()

            # Use the upload and poll SDK helper to upload the files, add them to the vector store,
            # and poll the status of the file batch for completion.
            file_batch = client.beta.vector_stores.file_batches.upload_and_poll(
                vector_store_id=vector_store.id,
                files=[
                    (original_file.name, original_file_stream),
                    (student_file.name, student_file_stream)
                ]
            )
            # attaching to thread
            prompt = f'''
            Please review the student exam file named {student_file.name} using the uploaded book {original_file.name} as the reference material to identify any errors. Additionally, please provide the student's final score and indicate the number of questions they answered incorrectly. Provide the answer in a structured JSON format as specified below.

            Instructions:
            1. Review the student exam file titled {student_file.name} using the uploaded book as a reference to detect errors.
            2. Determine the student's final score and specify the number of questions they answered incorrectly.
            3. Respond in a structured manner to improve clarity. Avoid manually evaluating each question individually.
            
            Response Format should be this json like format please :
            {{
            "review": {{
                "student_name": "[Student Name]",
                "book_name":[Book name]
                "results": {{
                "final_score": [Final Score],
                "incorrect_answers_count": [Number of Incorrect Answers],
                "incorrect_answers_details": [
                    {{
                    "question_number": [Question Number],
                    "reason": "[Reason for Incorrect Answer]",
                    "correct_answer": "[Correct Answer]"
                    }}
                    // Repeat for each incorrect answer
                ]
                }}
            }}
            }}

            please dont add intro text or outro text only the response format as response 
            '''
            thread = client.beta.threads.create(
                messages=[ { "role": "user", "content": prompt} ],
                tool_resources={
                    "file_search": {
                    "vector_store_ids": [vector_store.id]
                    }
                }
            )
            # run
            run = client.beta.threads.runs.create_and_poll(
                thread_id=thread.id, assistant_id=assistant2
            )
            # Poll until the run status is completed
            while run.status != "completed":
                time.sleep(2)  # Adding a delay to avoid rapid polling
                run = client.beta.threads.runs.retrieve(run.id,thread_id=thread.id)

            messages = list(client.beta.threads.messages.list(thread_id=thread.id, run_id=run.id))
            if messages:
                message_content = messages[0].content[0].text
                annotations = message_content.annotations
                citations = []
                for index, annotation in enumerate(annotations):
                    message_content.value = message_content.value.replace(annotation.text, f"[{index}]")
                    if file_citation := getattr(annotation, "file_citation", None):
                        cited_file = client.files.retrieve(file_citation.file_id)
                        citations.append(f"[{index}] {cited_file.filename}")
                # structured response
                structured_data = extract_details_from_string(message_content.value)
                save_review_to_db(structured_data,self.request.user)
                # Extract and structure the response
                # Return the status and file counts of the batch to see the result of this operation.
                return Response({
                    'status': file_batch.status,
                    'file_counts': file_batch.file_counts,
                    'response':structured_data,
                }, status=status.HTTP_200_OK)
            else :
                return Response({
                    'status': file_batch.status,
                    'file_counts': file_batch.file_counts,
                    'response':"No response"
                }, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class RecordStudentMarkView(generics.ListCreateAPIView):
    serializer_class = StudentMarkSerializer

    def get_queryset(self):
        return StudentReview.objects.filter(user=self.request.user)

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)

class DownloadStudentMarksCsvView(APIView):
    serializer_class = StudentMarkSerializer

    def get(self, request):
        marks = StudentReview.objects.filter(user=request.user)
        
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="student_marks.csv"'

        writer = csv.writer(response)
        writer.writerow(['Student Name', 'Book Name', 'Final Score', 'Incorrect Answers', 'Question Number', 'Status', 'Correct Version'])
        
        for mark in marks:
            detailed_reviews = QuestionDetail.objects.filter(student_review=mark)
            for detail in detailed_reviews:
                writer.writerow([
                    mark.student_name,
                    mark.book_name,
                    mark.final_score,
                    mark.incorrect_answers_count,
                    detail.question_number,
                    detail.status,
                    detail.correct_version
                ])

        return response
    
class StudentMarkDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = StudentMarkSerializer

    def get_queryset(self):
        return StudentReview.objects.filter(user=self.request.user)
    
class UserListCreateView(generics.ListCreateAPIView):
    queryset = User.objects.all()
    serializer_class = CustomUserSerializer
    permission_classes = [IsAdminUser]

class UserRetrieveUpdateDestroyView(generics.RetrieveUpdateDestroyAPIView):
    queryset = User.objects.all()
    serializer_class = CustomUserSerializer
    permission_classes = [IsAdminUser]

















